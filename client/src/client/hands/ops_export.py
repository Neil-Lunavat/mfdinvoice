"""The month's files, for the person: a picture of a signed invoice, and the month as a ZIP with an Excel register.

Both happen on this PC and nothing about them goes anywhere. The invoice picture is the signed PDF's first page,
drawn here; the export is every signed PDF of the month the window shows plus a register of its invoices, written into
the person's Downloads folder.
"""

from __future__ import annotations

import base64
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from client.brand import NAME

REGISTER = ["Registrar", "Fund house", "Fund house GSTIN", "Invoice number", "Date", "Taxable", "CGST", "SGST",
            "IGST", "Total", "Status", "The registrar's word", "Number in Tally", "Signed PDF"]


def first_page(pdf: Path, width: int = 900) -> str:
    """The first page as a PNG data URL, for the window's preview."""
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf))
    try:
        doc.init_forms()                           # a token's visible mark is a form field, drawn only with this
        page = doc[0]
        scale = width / page.get_width()
        image = page.render(scale=scale).to_pil()
    finally:
        doc.close()
    buf = io.BytesIO()
    image.save(buf, "PNG", optimize=True)
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def downloads(workspace: Path) -> Path:
    """The person's Downloads folder, or the app's own export folder when there is none."""
    home = Path.home() / "Downloads"
    return home if home.is_dir() else workspace / "export"


def export(month: dict, files: dict[str, Path | None], folder: Path) -> Path:
    """`MFDInvoice September 2026.zip`: the signed PDFs, and `Register September 2026.xlsx` listing every invoice."""
    folder.mkdir(parents=True, exist_ok=True)
    out = folder / f"{NAME} {month['label']}.zip"
    n = 2
    while out.exists():
        out = folder / f"{NAME} {month['label']} ({n}).zip"
        n += 1

    book = Workbook()
    sheet = book.active
    sheet.title = "Register"
    sheet.append(REGISTER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    names: dict[str, str] = {}
    for i in sorted(month["invoices"], key=lambda x: (x["registrar"], x["amc"].lower())):
        pdf = files.get(i["key"])
        folder_in_zip = "CAMS" if i["registrar"] == "CAMS" else "KFintech"
        if pdf and pdf.exists():
            names[i["key"]] = f"{folder_in_zip}/{pdf.name}"
        total = round(i["taxable"] + i["cgst"] + i["sgst"] + i["igst"], 2)
        in_tally = i.get("books") or i.get("tally") or ""
        sheet.append([folder_in_zip, i["amc"], i.get("gstin") or "", i["number"] or i["key"], i["date"], i["taxable"],
                      i["cgst"], i["sgst"], i["igst"], total, i["status"], i["said"],
                      "" if in_tally == "in" else in_tally, names.get(i["key"], "")])
    for column, width in zip("ABCDEFGHIJKLMN", (10, 28, 18, 22, 12, 12, 10, 10, 10, 12, 18, 26, 16, 40)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, min_col=6, max_col=10):
        for cell in row:
            cell.number_format = "#,##,##0.00"
    register = io.BytesIO()
    book.save(register)

    tmp = out.with_suffix(".part")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"Register {month['label']}.xlsx", register.getvalue())
        for key, name in names.items():
            z.write(files[key], name)
    tmp.replace(out)
    return out
