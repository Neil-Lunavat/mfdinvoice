"""Spreadsheets, archives and a PDF's text, read and written on this PC.

CAMS's upload page says, in its own words, "Use Prefilled Template … Do not alter any core system values", and the
emailed .xls is that template. `sheet_fill` reads that template and writes a new workbook from its cells, changing only the
cells it is told to, and every text cell is written as text (format "@"). That is what stops a payment month of
`092026` quietly becoming the number `92026`.
"""

from __future__ import annotations

import warnings
import zipfile
from pathlib import Path

import openpyxl
import pdfplumber
import xlrd

from client.automation.page import Changed, Stop

# KFintech's spreadsheets carry no default style, and openpyxl says so on every one. It changes nothing read.
warnings.filterwarnings("ignore", message="Workbook contains no default style")


def _cell_text(value) -> str:
    """One cell as text, the way the file means it: an integer stored as a number comes back as `26`, not `26.0`,
    because these are compared against what the portal shows on screen."""
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip() if isinstance(value, str) else str(value)


def sheet_read(path: Path, sheet: str = "", max_rows: int = 5000) -> list[list[str]]:
    """Every row of a sheet, each cell as text, exactly as it is. What the cells mean is decided by whoever asked."""
    rows: list[list[str]] = []
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path)
        sh = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
        for r in range(min(sh.nrows, max_rows + 1)):
            rows.append([_cell_text(sh.cell_value(r, c)) for c in range(sh.ncols)])
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > max_rows:
                    break
                rows.append([_cell_text(v) for v in row])
        finally:
            wb.close()
    else:
        raise Changed(f"{path.name} is not a spreadsheet")
    return rows


def zip_extract(path: Path, folder: Path) -> list[Path]:
    """Unpack an archive flat into a folder. A file already there with the same contents is left alone; a different
    one is replaced, so nothing is ever there twice; one held open by another program stops with its name."""
    folder.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            name = Path(info.filename).name                  # never trust a path inside an archive
            if not name:
                continue
            dest = folder / name
            data = z.read(info)
            try:
                same = dest.exists() and dest.read_bytes() == data
            except OSError:
                same = False
            if not same:
                try:
                    dest.write_bytes(data)
                except PermissionError:
                    raise Stop("wrong_files", f"{name} is open in another program", "Close it, then Run again.")
            written.append(dest)
    return written


def make_zip(out: Path, members: list[Path]) -> Path:
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for p in members:
            z.write(p, p.name)                               # flat: CAMS matches on the file name alone
    return out


def sheet_fill(src: Path, out: Path, edits: list[dict], keep: set[int] | None = None) -> Path:
    """Read a registrar's own template and write a new workbook from its cells, changing only those named in `edits`
    ({row, column, value}; row 0 is the first under the headings). Numbers stay numbers; text, including every edited
    value, is written as text. With `keep` (row numbers, counted as in `edits`), every other row is left out."""
    out.parent.mkdir(parents=True, exist_ok=True)
    if src.suffix.lower() == ".xls":
        book = xlrd.open_workbook(src)
        sh = book.sheet_by_index(0)
        headers = [str(sh.cell_value(0, c)).strip() for c in range(sh.ncols)]
        grid = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
        title = sh.name
    else:
        wb_in = openpyxl.load_workbook(src, data_only=True)
        ws_in = wb_in.worksheets[0]
        grid = [list(r) for r in ws_in.iter_rows(values_only=True)]
        headers = [str(h).strip() if h is not None else "" for h in (grid[0] if grid else [])]
        title = ws_in.title

    where = {h: i for i, h in enumerate(headers)}
    for e in edits:
        col = where.get(e["column"])
        if col is None:
            raise Changed(f"the template has no column called {e['column']!r} (it has: {headers})")
        r = e["row"] + 1                                     # 0 is the header row
        if not 1 <= r < len(grid):
            raise Changed(f"row {e['row']} is not in this template")
        grid[r][col] = e["value"]                            # a string: openpyxl writes it as text
    if keep is not None:
        grid = grid[:1] + [row for i, row in enumerate(grid[1:]) if i in keep]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in grid:
        ws.append(row)
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.number_format = "@"
    wb.save(out)
    return out


def text_layer(path: Path) -> dict:
    """Every word of a PDF with where it sits, in points from the bottom-left corner: {items, page_w, page_h}."""
    items, page_w, page_h = [], 0.0, 0.0
    with pdfplumber.open(path) as pdf:
        for n, page in enumerate(pdf.pages, start=1):
            page_w, page_h = float(page.width), float(page.height)
            for w in page.extract_words():
                items.append({
                    "page": n, "text": w["text"],
                    "x": round(float(w["x0"]), 2), "y": round(page_h - float(w["bottom"]), 2),
                    "w": round(float(w["x1"]) - float(w["x0"]), 2),
                    "h": round(float(w["bottom"]) - float(w["top"]), 2),
                })
    return {"items": items, "page_w": page_w, "page_h": page_h}
